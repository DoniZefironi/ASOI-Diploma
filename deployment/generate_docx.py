"""Конвертация Инструкция_по_развертыванию.md -> Инструкция_по_развертыванию.docx"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = Path(__file__).parent / "Инструкция_по_развертыванию.md"
DST = Path(__file__).parent / "Инструкция_по_развертыванию.docx"

URL_RE = re.compile(r'https?://[^\s\)\]\*`]+')
TOKEN_RE = re.compile(r'(\[[^\]]+\]\([^\)]+\)|\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|https?://[^\s\)\]\*`]+)')


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Consolas")
    rFonts.set(qn("w:hAnsi"), "Consolas")
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    rPr.append(sz)

    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_runs(paragraph, text):
    """Парсит **bold**, `code` и голые ссылки внутри строки."""
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("[") and "](" in token:
            link_text = token[1:token.index("](")]
            r = paragraph.add_run(link_text)
            r.italic = True
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`") and token.endswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        elif token.startswith("http"):
            url = token
            trailing = ""
            while url and url[-1] in ".,;:)":
                trailing = url[-1] + trailing
                url = url[:-1]
            add_hyperlink(paragraph, url, url)
            if trailing:
                paragraph.add_run(trailing)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def shade_paragraph(paragraph, color="EDEDED"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    pPr.append(shd)


def add_code_block(doc, code_lines, indent_cm=0.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent_cm)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    shade_paragraph(p)
    for i, line in enumerate(code_lines):
        if i > 0:
            p.add_run().add_break()
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
    return p


def add_table(doc, rows):
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            for run in cp.runs:
                run.text = ""
            add_runs(cp, cell_text)
            for run in cp.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()
    return table


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = 'TOC \\o "1-2" \\h \\z \\u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Нажмите Ctrl+A, затем F9, чтобы обновить содержание."
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")

    r_element = run._r
    r_element.append(fldChar1)
    r_element.append(instrText)
    r_element.append(fldChar2)
    r_element.append(t)
    r_element.append(fldChar3)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    for h in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[h]
        st.font.color.rgb = RGBColor(0x1F, 0x36, 0x5C)
        st.font.name = "Calibri"

    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(11.5)

    i = 0
    n = len(lines)
    skip_toc_list = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Заголовки
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped[2:])
            run.bold = True
            run.font.size = Pt(18)
            i += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:]
            doc.add_heading(text, level=1)
            skip_toc_list = (text == "Содержание")
            if skip_toc_list:
                add_toc(doc)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue

        # Горизонтальная линия / пропуск ручного содержания
        if stripped == "---":
            skip_toc_list = False
            i += 1
            continue

        if skip_toc_list:
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Блок кода
        if stripped.startswith("```"):
            indent = len(line) - len(line.lstrip())
            code_lines = []
            i += 1
            while i < n and lines[i].strip() != "```":
                code_lines.append(lines[i][indent:] if len(lines[i]) >= indent else lines[i].lstrip())
                i += 1
            i += 1  # пропустить закрывающий ```
            add_code_block(doc, code_lines, indent_cm=0.5 + (indent / 4) * 0.5)
            continue

        # Таблица
        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                row_line = lines[i].strip().strip("|")
                cells = [c.strip() for c in row_line.split("|")]
                if not re.match(r'^[\s\-:|]+$', lines[i].strip()):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        # Цитата (blockquote), может занимать несколько строк
        if stripped.startswith(">"):
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_runs(p, " ".join(quote_lines))
            for run in p.runs:
                run.italic = True
            continue

        # Нумерованный список верхнего уровня: "N. текст"
        m_num = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m_num:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m_num.group(2))
            i += 1
            # Захват строк-продолжений (с отступом >= 3 пробелов, простой текст)
            while i < n:
                nxt = lines[i]
                nxt_stripped = nxt.strip()
                if not nxt_stripped:
                    break
                indent = len(nxt) - len(nxt.lstrip())
                if indent >= 3 and not nxt_stripped.startswith(("```", "-", ">", "|", "#")) \
                        and not re.match(r'^\d+\.\s', nxt_stripped):
                    p.add_run(" ")
                    add_runs(p, nxt_stripped)
                    i += 1
                else:
                    break
            continue

        # Маркированный список ("- текст"), включая вложенные (с отступом)
        m_bul = re.match(r'^(\s*)-\s+(.*)$', line)
        if m_bul:
            indent = len(m_bul.group(1))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_runs(p, m_bul.group(2))
            i += 1
            while i < n:
                nxt = lines[i]
                nxt_stripped = nxt.strip()
                if not nxt_stripped:
                    break
                nindent = len(nxt) - len(nxt.lstrip())
                if nindent >= 2 and not nxt_stripped.startswith(("```", "-", ">", "|", "#")) \
                        and not re.match(r'^\d+\.\s', nxt_stripped):
                    p.add_run(" ")
                    add_runs(p, nxt_stripped)
                    i += 1
                else:
                    break
            continue

        # Обычный абзац (может занимать несколько строк до пустой строки)
        para_lines = [stripped]
        i += 1
        while i < n and lines[i].strip() and not lines[i].strip().startswith(("#", "|", ">", "```", "-")) \
                and not re.match(r'^\d+\.\s', lines[i].strip()):
            para_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(para_lines))

    doc.save(str(DST))
    print(f"Saved: {DST}")


if __name__ == "__main__":
    main()
